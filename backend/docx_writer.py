"""Apply validated operation plans to DOCX files."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def apply_operation_plan(template_path: str | Path, plan: dict[str, Any], output_path: str | Path) -> dict[str, Any]:
    document = Document(str(template_path))
    report: list[dict[str, Any]] = []

    for operation in plan.get("operations", []):
        if not isinstance(operation, dict) or operation.get("op") != "set_cell_text":
            continue
        target = operation.get("target", {})
        try:
            table_index = int(target.get("table"))
            row_index = int(target.get("row"))
            cell_index = int(target.get("cell"))
            row = document.tables[table_index].rows[row_index]
            cell = row.cells[cell_index]
        except Exception as exc:  # noqa: BLE001
            report.append({"field": operation.get("field"), "status": "skipped", "message": str(exc)})
            continue
        value = str(operation.get("value", ""))
        field = str(operation.get("field") or "")
        set_row_cant_split(row)
        write_cell_text(cell, value, compact=is_compact_field(field))
        report.append(
            {
                "field": field,
                "status": "written",
                "value": value,
                "target": target,
            }
        )

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    return {"docx": str(output), "report": report}


def write_cell_text(cell: Any, value: str, compact: bool = False) -> None:
    paragraph = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
    if paragraph.runs:
        paragraph.runs[0].text = value
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(value)
    if compact:
        apply_compact_paragraph(paragraph)
    for extra in cell.paragraphs[1:]:
        for run in extra.runs:
            run.text = ""
        if compact:
            apply_compact_paragraph(extra)


def is_compact_field(field: str) -> bool:
    return field.split("[", 1)[0] in {"careers", "certificates", "languages", "military"}


def apply_compact_paragraph(paragraph: Any) -> None:
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    for run in paragraph.runs:
        run.font.size = Pt(8.5)


def set_row_cant_split(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    for existing in tr_pr.findall(qn("w:cantSplit")):
        tr_pr.remove(existing)
    tr_pr.append(OxmlElement("w:cantSplit"))
